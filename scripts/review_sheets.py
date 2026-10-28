"""Review aids for native-speaker / human checks: contact sheets, mark-up workbooks, readable transcripts.

  sequences    Bloom-VIST picture stories (true order, left to right) + workbook
  near_twins   edited near-twin image sets (original + 5 edits, edit named) + workbook
  transcripts  a run's dialogues as Markdown (source text and what the partner received)

Writes to data/annotation/review/. Usage: python scripts/review_sheets.py [sequences near_twins transcripts]
  [--run dry2]
"""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw

from mtcomm.data import SEQUENCE_MANIFEST
from mtcomm.paths import DATA_ROOT, OUTPUT_ROOT

OUT = DATA_ROOT / "annotation" / "review"
THUMB = 220


def _font(size):
    from PIL import ImageFont

    p = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    return ImageFont.truetype(p, size) if Path(p).exists() else ImageFont.load_default()


def _thumb(path):
    im = Image.open(path).convert("RGB")
    im.thumbnail((THUMB, THUMB))
    return im


def _sheet(rows, path, label_h=36):
    """rows: [(title, [(image_path, caption)])] -> one PNG, one row per entry."""
    ncol = max(len(r[1]) for r in rows)
    w, h = 10 + ncol * (THUMB + 10), sum(THUMB + label_h + 30 for _ in rows) + 10
    sheet = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(sheet)
    f, fb = _font(12), _font(15)
    y = 10
    for title, cells in rows:
        d.text((10, y), title, fill="black", font=fb)
        y += 22
        for k, (p, cap) in enumerate(cells):
            x = 10 + k * (THUMB + 10)
            sheet.paste(_thumb(p), (x, y))
            d.text((x, y + THUMB + 2), cap[:34], fill=(60, 60, 60), font=f)
            d.text((x, y + THUMB + 16), cap[34:68], fill=(60, 60, 60), font=f)
        y += THUMB + label_h + 8
    sheet.save(path)


def _workbook(path, header, rows, instructions):
    wb = Workbook()
    ws = wb.active
    ws.append([instructions])
    ws["A1"].alignment = Alignment(wrap_text=True)
    ws.append(header)
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col, width in zip("ABCDEFGH", [16, 30, 14, 50, 22, 22, 30, 30]):
        ws.column_dimensions[col].width = width
    wb.save(path)


def sequences():
    m = json.load(open(SEQUENCE_MANIFEST))  # the screened set (v2) once it exists
    img = DATA_ROOT / "derived" / "bloom_vist" / "images"
    stories = m["stories"]
    for page in range(0, len(stories), 10):
        rows = [(f"S{page + i + 1:02d}  {s['title']}  ({s['licence']})",
                 [(img / f, f"picture {k + 1}") for k, f in enumerate(s["images"])])
                for i, s in enumerate(stories[page:page + 10])]
        _sheet(rows, OUT / f"sequence_stories_{page // 10 + 1}.png")
    _workbook(OUT / "sequence_stories.xlsx",
              ["story", "title", "licence", "story_id", "order clear from pictures alone? (y/n)",
               "usable (no offensive/odd content)? (y/n)", "note"],
              [[f"S{i + 1:02d}", s["title"], s["licence"], s["story_id"], "", "", ""] for i, s in enumerate(stories)],
              "Pictures are shown left to right in their true order (sequence_stories_*.png). Column E: could someone "
              "put them in order without being told? (y = guessable, which makes the game test less communication.) "
              "Column F: anything that should not be used.")


def near_twins():
    m = json.load(open(DATA_ROOT / "derived" / "perturb" / "full60" / "manifest_perturb_full60.json"))
    base = Path(m["image_dir"])
    if not base.is_absolute():
        base = DATA_ROOT / base
    ext = m.get("image_ext", ".png")
    items, wb_rows = m["items"], []
    for page in range(0, len(items), 8):
        rows = []
        for it in items[page:page + 8]:
            cells = [(base / f"{it['image_key']}{ext}", "ORIGINAL (correct)")]
            for key in it["distractors"]["hard"]:
                diff = it["differences"][key]
                cells.append((base / f"{key}{ext}", f"{diff['attribute']}: {diff['difference']}"))
                wb_rows.append([it["item_id"], key, diff["attribute"], diff["difference"], "", "", ""])
            rows.append((it["item_id"], cells))
        _sheet(rows, OUT / f"near_twins_{page // 8 + 1}.png")
    _workbook(OUT / "near_twins.xlsx",
              ["set", "edit", "attribute", "intended difference", "difference visible? (y/n)",
               "describable in a few words? (y/n)", "note"], wb_rows,
              "Each row of near_twins_*.png is one set: the original, then the 5 edited versions with the intended "
              "change written under each. Mark edits whose change you cannot see, or that look broken.")


def transcripts(run):
    recs = [json.loads(line) for line in open(OUTPUT_ROOT / run / "records.jsonl")]
    lines = [f"# Dialogue transcripts: run `{run}`\n",
             "Each message: what the player wrote, then what the partner received (after the translator). "
             "Ceiling = same language, no translator; prior = no messages.\n"]
    for r in sorted(recs, key=lambda r: (r["task"], r["episode"])):
        if r["status"] != "done":
            continue
        lines.append(f"\n## {r['episode']}\n\nroles: {r['roles']} · score: correct={r['score'].get('correct')}, "
                     f"guess={r['score'].get('guess')}, gold={r['score'].get('gold')}\n")
        for t in r["turns"]:
            if t["kind"] == "message":
                lines.append(f"- **{t['role']}** wrote: {t['source']}")
                if t["delivered"] != t["source"]:
                    lines.append(f"  - received: {t['delivered']}")
            else:
                lines.append(f"- **{t['role']} answered:** {t['source'][:200]}")
    (OUT / f"transcripts_{run}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("what", nargs="*", default=["sequences", "near_twins", "transcripts"])
    ap.add_argument("--run", default="dry2")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    for w in args.what:
        transcripts(args.run) if w == "transcripts" else globals()[w]()
        print("done:", w)
    print(OUT)


if __name__ == "__main__":
    main()
